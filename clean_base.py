import pandas as pd
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

df = pd.read_excel("knowledge_base.xlsx", header=1)

# الكلمات المشوّهة (من المشكلة)
bad_patterns = [
    "استشنورة", "المعلوإيشت", "اكتشنول",
    "استشنو", "شنورة", "إيشت",
    "شنوذا", "تحهذا", "استرهذا",
    "هذاد", "هذاذه"
]

# كلمات مشوّهة تانية (لو ظهرت)
generic_bad = [
    "شنوذا", "إيشي", "هذاذي", "شنوة", "إيشة"
]

bad_keywords = bad_patterns + generic_bad

# فلترة الصفوف
def is_bad(question):
    q = str(question)
    for bad in bad_keywords:
        if bad in q:
            return True
    return False

bad_rows = df[df["السؤال"].apply(is_bad)]
print(f"عدد الصفوف المشوّهة: {len(bad_rows)}")

if len(bad_rows) > 0:
    print("\nأمثلة:")
    for i, row in bad_rows.head(10).iterrows():
        print(f"- {row['السؤال']}")

# حذف الصفوف المشوّهة
df_clean = df[~df["السؤال"].apply(is_bad)]
print(f"\nعدد الصفوف قبل التنظيف: {len(df)}")
print(f"عدد الصفوف بعد التنظيف: {len(df_clean)}")

# كتابة الملف
title = "الأسئلة والأجوبة الخاصة بكلية علوم الحاسوب وتقانة المعلومات"

df_clean.to_excel("knowledge_base.xlsx", index=False, header=True, startrow=1, sheet_name="الأسئلة والأجوبة")

# التنسيق
wb = load_workbook("knowledge_base.xlsx")
ws = wb.active

navy = "1F2A5C"
red = "C1272D"
light_gray = "F5F3EE"
white = "FFFFFF"

title_font = Font(name="Cairo", size=16, bold=True, color=white)
title_fill = PatternFill("solid", fgColor=navy)
title_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

header_font = Font(name="Cairo", size=13, bold=True, color=white)
header_fill = PatternFill("solid", fgColor=red)
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

body_font = Font(name="Cairo", size=11)
body_align = Alignment(horizontal="right", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

ws.merge_cells("A1:B1")
ws["A1"] = title
ws["A1"].font = title_font
ws["A1"].fill = title_fill
ws["A1"].alignment = title_align
ws.row_dimensions[1].height = 40

for col in ["A", "B"]:
    cell = ws[f"{col}2"]
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border
ws.row_dimensions[2].height = 30

last_row = ws.max_row
for row in range(3, last_row + 1):
    cell_a = ws[f"A{row}"]
    cell_a.font = body_font
    cell_a.alignment = body_align
    cell_a.border = thin_border

    cell_b = ws[f"B{row}"]
    cell_b.font = body_font
    cell_b.alignment = body_align
    cell_b.border = thin_border

    if row % 2 == 0:
        cell_a.fill = PatternFill("solid", fgColor=light_gray)
        cell_b.fill = PatternFill("solid", fgColor=light_gray)

ws.column_dimensions["A"].width = 55
ws.column_dimensions["B"].width = 90
ws.freeze_panes = "A3"

for row in range(3, last_row + 1):
    ws.row_dimensions[row].height = 50

wb.save("knowledge_base.xlsx")
print("\nتم حفظ الملف بنجاح")
